"""Resume file parsing: TXT / DOCX / PDF, plus failure handling."""

from __future__ import annotations

import docx
import pytest

from app.parsing import ResumeParseError, extract_text
from app.services.demo_data import _write_broken_pdf, _write_docx, _write_simple_pdf


def test_txt_parsing(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text("Ravi Kumar\nravi@example.com\nPython, SQL, Git\n" * 3, encoding="utf-8")
    text = extract_text(path)
    assert "Ravi Kumar" in text
    assert "ravi@example.com" in text


def test_txt_parsing_cp1252_fallback(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_bytes("José Fernandes\nAccounts and Tally experience here.\n".encode("cp1252"))
    text = extract_text(path)
    assert "Fernandes" in text


def test_docx_parsing(tmp_path):
    path = tmp_path / "resume.docx"
    _write_docx(path, "Anita Rao\nanita@example.com\n\nSKILLS\nPython, SQL, Git\n\nPROJECTS\nPayroll tool using Python and SQL\n")
    text = extract_text(path)
    assert "Anita Rao" in text
    assert "Payroll tool" in text


def test_docx_corrupted_file(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"PK\x03\x04this is not really a docx file at all")
    with pytest.raises(ResumeParseError) as excinfo:
        extract_text(path)
    assert "corrupted" in excinfo.value.reason.lower() or "Word" in excinfo.value.reason


def test_pdf_parsing(tmp_path):
    path = tmp_path / "resume.pdf"
    _write_simple_pdf(
        path,
        [
            "Nikhil Reddy",
            "Email: nikhil@example.com  Phone: +91 98250 12121",
            "EDUCATION: BCA, Osmania University, CGPA 8.0/10",
            "SKILLS: Python, SQL, Git",
            "PROJECTS: Student portal using Flask and SQL",
        ],
    )
    text = extract_text(path)
    assert "Nikhil Reddy" in text
    assert "Python, SQL, Git" in text


def test_broken_pdf_fails_with_clean_reason(tmp_path):
    path = tmp_path / "broken.pdf"
    _write_broken_pdf(path)
    with pytest.raises(ResumeParseError) as excinfo:
        extract_text(path)
    assert "PDF" in excinfo.value.reason


def test_too_short_text_is_rejected(tmp_path):
    path = tmp_path / "tiny.txt"
    path.write_text("hi", encoding="utf-8")
    with pytest.raises(ResumeParseError):
        extract_text(path)


def test_missing_file_is_rejected(tmp_path):
    with pytest.raises(ResumeParseError):
        extract_text(tmp_path / "missing.pdf")


def test_docx_tables_are_read(tmp_path):
    document = docx.Document()
    document.add_paragraph("Meera Nair")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python, SQL"
    path = tmp_path / "table_resume.docx"
    document.save(str(path))
    text = extract_text(path)
    assert "Meera Nair" in text
    assert "Python, SQL" in text
